"""The PO to the brand as an Excel file (PRD §4.1 step 4, §4.4.5).

The layout is the one the user approved, `docs/templates/PO Restock - Kahf -
MA5.xlsx`, made by `tools/gen_excel_templates.py`. This module is that
generator's PO part ported to live data: same cells, same styles, same column
widths. If the approved template changes, change both.

It is pure: header and lines in, bytes out, no database. That keeps it testable
without one, and keeps the router free of spreadsheet detail.

Differences from the template, all deliberate:
  * "Quantity requested" is the number Ops HQ saved, not the template's
    =MAX(0, fill up to - stock) formula: the PO quantities are frozen at
    *Simpan PO* and may have been changed by hand.
  * An optional "Note" header row (row 11) when Ops HQ wrote one; row 11 is
    blank in the template, so the rest of the sheet does not move.
  * A barcode is drawn only when it is a valid EAN-13 (or a UPC-A, which is an
    EAN-13 with a leading 0). Other codes are printed as digits only, because a
    drawn barcode that does not scan is worse than none.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

# ------------------------------------------------------------------ styles (as the template)
F = "Arial"
INK, MUTED, RED = "1B1D1C", "6B706D", "C0202D"
HEAD_FILL = PatternFill("solid", fgColor="1B1D1C")
TOTAL_FILL = PatternFill("solid", fgColor="F2F4F1")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="C6CBC5")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
QTY = '#,##0;(#,##0);"-"'

REQUESTS = ("Please: (1) make sure every unit carries the barcode in the Barcode column; if it "
            "differs or is missing, write the right barcode in the yellow column. (2) List the "
            "expiry date (ED) of each SKU on the Faktur. (3) Put this PO number on the Faktur "
            "and the Surat Jalan.")
HEADERS = ["No", "Brand SKU code", "Hiryu SKU code", "Barcode (EAN-13)", "Barcode", "Product",
           "Size", "Current stock", "Fill up to", "Quantity requested (pcs)",
           "Brand's barcode (fill if different or missing)", "Notes"]
WIDTHS = [5, 14, 14, 17, 26, 46, 9, 10, 10, 13, 24, 22]


def _font(**kw):
    return Font(name=F, color=kw.pop("color", INK), **kw)


def _cell(ws, r, c, v, fmt=None, bold=False, fill=None, color=INK, align=None):
    x = ws.cell(r, c, v)
    x.font = _font(size=10, bold=bold, color=color)
    x.border = BOX
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if align:
        x.alignment = Alignment(horizontal=align, vertical="center")
    return x


def _title_block(ws, title, rows):
    ws["A1"] = "NINJA VAN"
    ws["A1"].font = _font(bold=True, size=9, color=RED)
    ws["A2"] = title
    ws["A2"].font = _font(bold=True, size=16)
    r = 4
    for label, value in rows:
        ws.cell(r, 1, label).font = _font(size=9, color=MUTED)
        ws.cell(r, 2, value).font = _font(size=10, bold=True)
        r += 1
    return r + 1


def _header_row(ws, r, headers, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(r, i, h)
        c.font = _font(bold=True, size=9, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[r].height = 30
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _dmy(v) -> str:
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, str) and len(v) >= 10 and v[4] == "-":
        return f"{v[8:10]}/{v[5:7]}/{v[0:4]}"
    return v or ""


# ------------------------------------------------------------------ EAN-13
L = ["0001101", "0011001", "0010011", "0111101", "0100011", "0110001", "0101111", "0111011",
     "0110111", "0001011"]
G = ["0100111", "0110011", "0011011", "0100001", "0011101", "0111001", "0000101", "0010001",
     "0001001", "0010111"]
R = ["1110010", "1100110", "1101100", "1000010", "1011100", "1001110", "1010000", "1000100",
     "1001000", "1110100"]
PARITY = ["LLLLLL", "LLGLGG", "LLGGLG", "LLGGGL", "LGLLGG", "LGGLLG", "LGGGLL", "LGLGLG",
          "LGLGGL", "LGGLGL"]


def ean13(code: str | None) -> str | None:
    """The 13 digits to draw, or None when the code is not a valid EAN-13/UPC-A."""
    code = (code or "").strip()
    if len(code) == 12 and code.isdigit():
        code = "0" + code
    if len(code) != 13 or not code.isdigit():
        return None
    d = [int(c) for c in code]
    check = (10 - (sum(d[i] * (3 if i % 2 else 1) for i in range(12)) % 10)) % 10
    return code if check == d[12] else None


def ean13_png(code: str) -> io.BytesIO:
    d = [int(c) for c in code]
    bits = "101"
    for i, p in enumerate(PARITY[d[0]]):
        bits += (L if p == "L" else G)[d[i + 1]]
    bits += "01010"
    for i in range(7, 13):
        bits += R[d[i]]
    bits += "101"
    m, h = 2, 46
    img = Image.new("RGB", ((len(bits) + 18) * m, h + 16), "white")
    dr = ImageDraw.Draw(img)
    x = 9 * m
    for i, b in enumerate(bits):
        if b == "1":
            long_bar = i < 3 or 45 <= i < 50 or i >= 92
            dr.rectangle([x, 2, x + m - 1, h + (8 if long_bar else 0)], fill="black")
        x += m
    dr.text((1, h + 2), code[0], fill="black")
    dr.text((13 * m, h + 3), code[1:7], fill="black")
    dr.text((59 * m, h + 3), code[7:], fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


# ------------------------------------------------------------------ the PO
def build_po_xlsx(header: dict, lines: list[dict]) -> bytes:
    """Build the PO workbook.

    header: reference, po_date, po_to, po_brand_contact, po_deliver_to,
            po_requested_date, po_receiving_hours, po_created_by_name, po_note
    lines:  brand_sku_code, hiryu_sku_code, barcode (None = MISSING), name,
            size, current_stock, fill_to, qty, note
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "PO"
    requested = _dmy(header.get("po_requested_date"))
    hours = (header.get("po_receiving_hours") or "").strip()
    if hours:
        requested = f"{requested}, receiving {hours}" if requested else f"Receiving {hours}"
    _title_block(ws, "Purchase Order · Restock", [
        ("PO number", header.get("reference") or ""),
        ("PO date", _dmy(header.get("po_date"))),
        ("To (brand)", header.get("po_to") or ""),
        ("Brand contact", header.get("po_brand_contact") or ""),
        ("Deliver to", header.get("po_deliver_to") or ""),
        ("Requested delivery date", requested),
        ("Created by (Ops HQ)", header.get("po_created_by_name") or ""),
    ])
    if (header.get("po_note") or "").strip():
        ws.cell(11, 1, "Note").font = _font(size=9, color=MUTED)
        ws.cell(11, 2, header["po_note"].strip()).font = _font(size=10, bold=True)
    ws["A12"] = REQUESTS
    ws["A12"].font = _font(size=10, bold=True, color=RED)
    ws["A12"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A12:L12")
    ws.row_dimensions[12].height = 44

    start = 14
    _header_row(ws, start, HEADERS, WIDTHS)
    r = start + 1
    first = r
    for i, s in enumerate(lines, 1):
        digits = (s.get("barcode") or "").strip() or None
        drawable = ean13(digits)
        _cell(ws, r, 1, i, align="center")
        _cell(ws, r, 2, s.get("brand_sku_code"))
        _cell(ws, r, 3, s.get("hiryu_sku_code"))
        b = _cell(ws, r, 4, digits or "MISSING", color=INK if digits else RED, bold=not digits)
        b.number_format = "@"   # digits as text: Excel must not turn them into 8.99E+12
        _cell(ws, r, 5, None)
        _cell(ws, r, 6, s.get("name"))
        _cell(ws, r, 7, s.get("size"))
        _cell(ws, r, 8, s.get("current_stock"), QTY)
        _cell(ws, r, 9, s.get("fill_to"), QTY)
        _cell(ws, r, 10, int(s.get("qty") or 0), QTY, bold=True)
        k = _cell(ws, r, 11, None, fill=INPUT_FILL)
        k.number_format = "@"
        _cell(ws, r, 12, s.get("note"))
        ws.row_dimensions[r].height = 50
        if drawable:
            img = XLImage(ean13_png(drawable))
            img.width, img.height = 180, 58
            ws.add_image(img, f"E{r}")
        for c in range(1, 13):
            ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=c in (6, 12),
                                                horizontal="center" if c == 1 else None)
        r += 1
    last = max(first, r - 1)
    for c in range(1, 13):
        _cell(ws, r, c, None, fill=TOTAL_FILL)
    _cell(ws, r, 4, f'=COUNTIF(D{first}:D{last},"MISSING")&" SKU without a barcode"',
          bold=True, fill=TOTAL_FILL, color=RED)
    _cell(ws, r, 6, "Total", bold=True, fill=TOTAL_FILL)
    _cell(ws, r, 10, f"=SUM(J{first}:J{last})", QTY, bold=True, fill=TOTAL_FILL)
    r += 2
    ws.cell(r, 1, "Yellow cells are for the brand to fill. MISSING means the WMS has no "
                  "barcode for that SKU yet.").font = _font(size=9, color=MUTED)
    r += 2
    for label in ("Created by Ops HQ (name, date)",
                  "Confirmed by the brand (name, date, shipment AWB)"):
        ws.cell(r, 2, label).font = _font(size=9, color=MUTED)
        for col in range(2, 7):
            ws.cell(r + 2, col).border = Border(bottom=thin)
        r += 4
    ws.freeze_panes = ws.cell(start + 1, 4)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    wb.calculation.fullCalcOnLoad = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
